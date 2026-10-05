import openpyxl

wb = openpyxl.load_workbook("data/review/ai_pre_review.xlsx")
sheet = wb["Patch for builder"]
headers = [c for c in next(sheet.iter_rows(values_only=True))]
print("Headers in Patch for builder:", headers)

col_vals = set()
for row in sheet.iter_rows(min_row=2, values_only=True):
    for i, val in enumerate(row):
        if val in ("HOLD_DRAFT", "REJECT_SUGGESTED", "AI_PRE_REVIEWED", "DISCLOSED_UNSCORED"):
            print(f"Col {i} ({headers[i]}): {val}")
            col_vals.add(headers[i])
            break
print("Matching status columns:", col_vals)
